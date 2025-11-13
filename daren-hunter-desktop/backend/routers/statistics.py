from fastapi import APIRouter, Depends
from sqlalchemy.orm import Session
from sqlalchemy import func, and_
from database import get_db, SendLog, Template
from schemas import Statistics, TemplateStats, TimeSlotStats
from datetime import datetime, timedelta
from typing import List

router = APIRouter(prefix="/api/statistics", tags=["数据统计"])


@router.get("/overview", response_model=Statistics)
def get_statistics_overview(db: Session = Depends(get_db)):
    """获取统计概览"""
    now = datetime.now()
    today = now.date()
    week_ago = now - timedelta(days=7)
    month_ago = now - timedelta(days=30)

    # 今日统计
    today_sent = db.query(func.count(SendLog.id)).filter(
        and_(
            func.date(SendLog.send_time) == today,
            SendLog.status == "success"
        )
    ).scalar() or 0

    today_replied = db.query(func.count(SendLog.id)).filter(
        and_(
            func.date(SendLog.send_time) == today,
            SendLog.reply_status == "replied"
        )
    ).scalar() or 0

    # 本周统计
    week_sent = db.query(func.count(SendLog.id)).filter(
        and_(
            SendLog.send_time >= week_ago,
            SendLog.status == "success"
        )
    ).scalar() or 0

    week_replied = db.query(func.count(SendLog.id)).filter(
        and_(
            SendLog.send_time >= week_ago,
            SendLog.reply_status == "replied"
        )
    ).scalar() or 0

    # 本月统计
    month_sent = db.query(func.count(SendLog.id)).filter(
        and_(
            SendLog.send_time >= month_ago,
            SendLog.status == "success"
        )
    ).scalar() or 0

    month_replied = db.query(func.count(SendLog.id)).filter(
        and_(
            SendLog.send_time >= month_ago,
            SendLog.reply_status == "replied"
        )
    ).scalar() or 0

    return Statistics(
        today_sent=today_sent,
        today_replied=today_replied,
        today_reply_rate=round(today_replied / today_sent * 100, 2) if today_sent > 0 else 0,
        week_sent=week_sent,
        week_replied=week_replied,
        week_reply_rate=round(week_replied / week_sent * 100, 2) if week_sent > 0 else 0,
        month_sent=month_sent,
        month_replied=month_replied,
        month_reply_rate=round(month_replied / month_sent * 100, 2) if month_sent > 0 else 0
    )


@router.get("/templates", response_model=List[TemplateStats])
def get_template_statistics(db: Session = Depends(get_db)):
    """获取话术效果统计"""
    templates = db.query(Template).all()
    stats = []

    for template in templates:
        usage_count = template.usage_count or 0
        reply_count = template.reply_count or 0
        reply_rate = round(reply_count / usage_count * 100, 2) if usage_count > 0 else 0

        stats.append(TemplateStats(
            template_name=template.name,
            usage_count=usage_count,
            reply_count=reply_count,
            reply_rate=reply_rate
        ))

    # 按回复率排序
    stats.sort(key=lambda x: x.reply_rate, reverse=True)
    return stats


@router.get("/best-time", response_model=List[TimeSlotStats])
def get_best_time_statistics(db: Session = Depends(get_db)):
    """获取最佳发送时间统计"""
    # 定义时间段
    time_slots = [
        ("00:00-06:00", 0, 6),
        ("06:00-09:00", 6, 9),
        ("09:00-12:00", 9, 12),
        ("12:00-14:00", 12, 14),
        ("14:00-17:00", 14, 17),
        ("17:00-19:00", 17, 19),
        ("19:00-22:00", 19, 22),
        ("22:00-24:00", 22, 24),
    ]

    stats = []
    for slot_name, start_hour, end_hour in time_slots:
        # 查询该时间段的发送记录
        logs = db.query(SendLog).filter(
            and_(
                SendLog.status == "success",
                func.extract('hour', SendLog.send_time) >= start_hour,
                func.extract('hour', SendLog.send_time) < end_hour
            )
        ).all()

        total = len(logs)
        replied = sum(1 for log in logs if log.reply_status == "replied")
        reply_rate = round(replied / total * 100, 2) if total > 0 else 0

        if total > 0:  # 只返回有数据的时间段
            stats.append(TimeSlotStats(
                time_slot=slot_name,
                reply_count=replied,
                reply_rate=reply_rate
            ))

    # 按回复率排序
    stats.sort(key=lambda x: x.reply_rate, reverse=True)
    return stats


@router.get("/export-report")
def export_statistics_report(db: Session = Depends(get_db)):
    """导出统计报表为Excel"""
    import openpyxl
    from io import BytesIO
    from fastapi.responses import StreamingResponse

    workbook = openpyxl.Workbook()

    # 概览统计
    overview_sheet = workbook.active
    overview_sheet.title = "数据概览"
    overview_stats = get_statistics_overview(db)

    overview_sheet.append(["指标", "今日", "本周", "本月"])
    overview_sheet.append([
        "发送数量",
        overview_stats.today_sent,
        overview_stats.week_sent,
        overview_stats.month_sent
    ])
    overview_sheet.append([
        "回复数量",
        overview_stats.today_replied,
        overview_stats.week_replied,
        overview_stats.month_replied
    ])
    overview_sheet.append([
        "回复率",
        f"{overview_stats.today_reply_rate}%",
        f"{overview_stats.week_reply_rate}%",
        f"{overview_stats.month_reply_rate}%"
    ])

    # 话术统计
    template_sheet = workbook.create_sheet("话术效果")
    template_sheet.append(["话术名称", "使用次数", "回复次数", "回复率"])
    template_stats = get_template_statistics(db)
    for stat in template_stats:
        template_sheet.append([
            stat.template_name,
            stat.usage_count,
            stat.reply_count,
            f"{stat.reply_rate}%"
        ])

    # 时间段统计
    time_sheet = workbook.create_sheet("最佳时间")
    time_sheet.append(["时间段", "回复次数", "回复率"])
    time_stats = get_best_time_statistics(db)
    for stat in time_stats:
        time_sheet.append([
            stat.time_slot,
            stat.reply_count,
            f"{stat.reply_rate}%"
        ])

    # 保存到内存
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=statistics_report_{datetime.now().strftime('%Y%m%d')}.xlsx"}
    )
