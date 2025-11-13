from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from sqlalchemy.orm import Session
from typing import List, Optional
from database import get_db, Target
from schemas import TargetCreate, TargetUpdate, TargetResponse, TargetListResponse, ImportResult
import openpyxl
from io import BytesIO

router = APIRouter(prefix="/api/targets", tags=["达人管理"])


@router.post("/", response_model=TargetResponse)
def create_target(target: TargetCreate, db: Session = Depends(get_db)):
    """创建新达人"""
    # 检查是否已存在
    existing = db.query(Target).filter(Target.douyin_id == target.douyin_id).first()
    if existing:
        raise HTTPException(status_code=400, detail="该抖音ID已存在")

    db_target = Target(**target.dict())
    db.add(db_target)
    db.commit()
    db.refresh(db_target)
    return db_target


@router.get("/", response_model=TargetListResponse)
def list_targets(
    skip: int = 0,
    limit: int = 100,
    status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    """获取达人列表"""
    query = db.query(Target)

    # 按状态筛选
    if status:
        query = query.filter(Target.status == status)

    # 搜索
    if search:
        query = query.filter(
            (Target.nickname.like(f"%{search}%")) |
            (Target.douyin_id.like(f"%{search}%"))
        )

    total = query.count()
    items = query.order_by(Target.created_at.desc()).offset(skip).limit(limit).all()

    return {"total": total, "items": items}


@router.get("/{target_id}", response_model=TargetResponse)
def get_target(target_id: int, db: Session = Depends(get_db)):
    """获取达人详情"""
    target = db.query(Target).filter(Target.id == target_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="达人不存在")
    return target


@router.put("/{target_id}", response_model=TargetResponse)
def update_target(target_id: int, target: TargetUpdate, db: Session = Depends(get_db)):
    """更新达人信息"""
    db_target = db.query(Target).filter(Target.id == target_id).first()
    if not db_target:
        raise HTTPException(status_code=404, detail="达人不存在")

    for key, value in target.dict(exclude_unset=True).items():
        setattr(db_target, key, value)

    db.commit()
    db.refresh(db_target)
    return db_target


@router.delete("/{target_id}")
def delete_target(target_id: int, db: Session = Depends(get_db)):
    """删除达人"""
    db_target = db.query(Target).filter(Target.id == target_id).first()
    if not db_target:
        raise HTTPException(status_code=404, detail="达人不存在")

    db.delete(db_target)
    db.commit()
    return {"message": "删除成功"}


@router.post("/batch-delete")
def batch_delete_targets(target_ids: List[int], db: Session = Depends(get_db)):
    """批量删除达人"""
    db.query(Target).filter(Target.id.in_(target_ids)).delete(synchronize_session=False)
    db.commit()
    return {"message": f"已删除 {len(target_ids)} 个达人"}


@router.post("/import-excel", response_model=ImportResult)
async def import_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """从Excel导入达人"""
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="请上传Excel文件")

    content = await file.read()
    workbook = openpyxl.load_workbook(BytesIO(content))
    sheet = workbook.active

    success = 0
    failed = 0
    duplicates = 0
    errors = []

    # 跳过表头
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):  # 跳过空行
            continue

        try:
            nickname = row[0]
            douyin_id = row[1]
            homepage_url = row[2] if len(row) > 2 else None
            fans_count = int(row[3]) if len(row) > 3 and row[3] else 0
            tags = row[4] if len(row) > 4 else None
            notes = row[5] if len(row) > 5 else None

            if not nickname or not douyin_id:
                errors.append(f"第{row_idx}行：昵称或抖音ID为空")
                failed += 1
                continue

            # 检查是否已存在
            existing = db.query(Target).filter(Target.douyin_id == douyin_id).first()
            if existing:
                duplicates += 1
                continue

            # 创建新记录
            target = Target(
                nickname=nickname,
                douyin_id=douyin_id,
                homepage_url=homepage_url,
                fans_count=fans_count,
                tags=tags,
                notes=notes,
                status="pending"
            )
            db.add(target)
            success += 1

        except Exception as e:
            errors.append(f"第{row_idx}行：{str(e)}")
            failed += 1

    db.commit()

    return ImportResult(
        success=success,
        failed=failed,
        duplicates=duplicates,
        total=success + failed + duplicates,
        errors=errors
    )


@router.get("/download/template")
def download_template():
    """下载Excel导入模板"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "达人列表"

    # 设置表头
    headers = ["昵称*", "抖音ID*", "主页链接", "粉丝数", "标签", "备注"]
    sheet.append(headers)

    # 添加示例数据
    sheet.append(["AI工坊", "xxx123", "https://www.douyin.com/user/xxx123", "28000", "AI视频,Sora", "优质作者"])

    # 保存到内存
    from io import BytesIO
    from fastapi.responses import StreamingResponse

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=daren_import_template.xlsx"}
    )


@router.post("/export")
def export_targets(target_ids: Optional[List[int]] = None, db: Session = Depends(get_db)):
    """导出达人数据为Excel"""
    query = db.query(Target)
    if target_ids:
        query = query.filter(Target.id.in_(target_ids))

    targets = query.all()

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "达人列表"

    # 设置表头
    headers = ["ID", "昵称", "抖音ID", "主页链接", "粉丝数", "标签", "状态", "备注", "创建时间"]
    sheet.append(headers)

    # 添加数据
    for target in targets:
        sheet.append([
            target.id,
            target.nickname,
            target.douyin_id,
            target.homepage_url,
            target.fans_count,
            target.tags,
            target.status,
            target.notes,
            target.created_at.strftime("%Y-%m-%d %H:%M:%S")
        ])

    # 保存到内存
    from io import BytesIO
    from fastapi.responses import StreamingResponse

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=daren_export.xlsx"}
    )
